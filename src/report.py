"""Build outputs/report.docx: split/validation approach, results, and the December chart."""
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches

from . import data

OUT = data.ROOT / "outputs"


def fold_chart(cv):
    d = cv[(cv.scope == "clean") & cv.model.isin(["baseline_rpm", "ridge", "lgb_no_market", "lgb_full"])]
    p = d.pivot(index="fold", columns="model", values="MAPE")
    ax = p.plot(marker="o", figsize=(7, 3.4))
    ax.set_ylabel("MAPE (%)"); ax.set_xlabel("Test month (trained on all earlier months)")
    ax.set_title("Time-based CV by fold (clean labels)"); ax.grid(alpha=.3)
    plt.tight_layout(); plt.savefig(OUT / "cv_by_fold.png", dpi=150); plt.close()


def main():
    cv = pd.read_csv(OUT / "cv_results.csv")
    summ = pd.read_csv(OUT / "cv_summary.csv")
    fold_chart(cv)
    doc = Document()
    doc.add_heading("Freight Rate Prediction: Approach and Validation", 0)

    doc.add_heading("1. Split and validation approach", 1)
    for t in [
        "validation.csv is a later period (Nov-Dec 2025) than train-test.csv (Jan-Oct 2025), so random K-fold "
        "would leak time. I used an expanding-window split: for each test month M in Jun-Oct, train on every "
        "month before M and score on M. This mimics the real task of predicting the next period from the past.",
        "Metrics: MAE, RMSE, MAPE and median APE. Spotter's metric is unknown, so I report several and picked "
        "the model that wins on all of them. Scores are reported on clean labels and on raw labels.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("2. Data-quality issues and fixes", 1)
    for t in [
        "About 1.4% of posted_rate values (677 rows) are corrupted. The corruption is near-symmetric: 340 rows "
        "sit about 3.5x above the normal rate per mile for their equipment and 337 sit about 3.6x below it, "
        "spread evenly across months and equipment types. Distance stays consistent with the coordinates on "
        "these rows, so the rate itself is wrong rather than the lane. They are detected with a deliberately "
        "low-capacity model (distance + equipment) whose log-residual exceeds 0.5, and removed from training "
        "only - never from the validation predictions, which must cover all 12,000 loads.",
        "Negative weights (292 train, 145 validation) are sign errors; their magnitudes match the positive "
        "distribution, so I take abs(). Missing weight (300 / 165) uses the equipment median plus an indicator.",
        "Missing market_index (374 / 249) is left as NaN with an indicator.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3. Model and results", 1)
    doc.add_paragraph(
        "Rate is driven by distance (corr 0.91) with a decreasing rate per mile, so I model log(posted_rate) "
        "with LightGBM using distance, equipment, weight, origin, destination, coordinates, circuity and "
        "day of week. Baselines: median rate per mile by equipment and distance bucket, and a Ridge model.")
    t = doc.add_table(rows=1, cols=6); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Scope", "Model", "MAE", "RMSE", "MAPE %", "MedAPE %"]):
        t.rows[0].cells[i].text = h
    for _, r in summ.iterrows():
        c = t.add_row().cells
        c[0].text, c[1].text = r["scope"], r["model"]
        for i, k in enumerate(["MAE", "RMSE", "MAPE", "MedAPE"]):
            c[2 + i].text = f"{r[k]:.2f}"
    doc.add_paragraph()
    doc.add_picture(str(OUT / "cv_by_fold.png"), width=Inches(6))
    doc.add_paragraph(
        "Market features (market_index, quote_signal) did not improve time-based CV, so the final model "
        "excludes them. I tested them raw (3.06% MAPE) and as four stationary transforms - deviation from "
        "the daily mean, level relative to a trailing 28-day mean, and 7d/28d momentum - and none beat the "
        "2.83% of the model without them; the best variant reached 2.95%.")
    doc.add_paragraph(
        "The per-fold pattern explains why. Raw market features win the early folds (2.84% vs 6.03% on June) "
        "and lose the late ones (4.62% vs 1.91% on September). That tracks the market regime: market_index "
        "climbs from 0.91 in January to 1.30 in May, then reverses and falls to 0.87 by September. While the "
        "trend continues the feature helps, but gradient-boosted trees cannot extrapolate beyond their "
        "trained range, so when the regime reverses the model is confidently wrong. The signal is real but "
        "unstable, and the instability costs more than the signal earns.")
    doc.add_paragraph(
        "The decisive folds are September and October, the closest analogues to the Nov-Dec prediction "
        "window, where excluding market features wins clearly (1.91% and 1.75% against 4.62% and 3.18%). "
        "market_index is also flat at roughly 0.92-0.95 across November and December, essentially unchanged "
        "from October, so there is no regime shift for a market feature to capture. The known failure mode "
        "is a holdout that spans a regime break, as June does.")
    doc.add_paragraph(
        "Excluding these features also lets one model serve both the 12,000 validation loads and the "
        "December chart, where they do not exist. The raw-label RMSE is dominated by the corrupted labels, "
        "which no model can predict.")

    doc.add_heading("4. December prediction", 1)
    doc.add_paragraph(
        "The fixed lane (Lexington to Fort Wayne, 360 mi, Dry Van, 32,000 lb) is scored with the same model, "
        "so only the date changes. The curve is a weekly cycle (about $814 to $832, mean $826, about $2.29/mi), "
        "consistent with the lane's own Jan-Oct history (mean $2.26/mi). There is no trend because the "
        "training data shows no reliable month-level effect once distance and equipment are accounted for.")
    doc.add_picture(str(data.ROOT / "scorer_results" / "candidate_december.png"), width=Inches(6.3))
    doc.save(OUT / "report.docx")


if __name__ == "__main__":
    main()
